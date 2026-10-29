import os
import csv
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Define the root directory and the combinations of interest
root_dir = '/sddata/projects/Conformal_Uncertainty_Quantification/Comparison_SL_SSL_PAVE/analysis/test2/SSL'
subdirectories = ['M36_Pipeline/M36_Norms']
# subdirectories = ['M36_Pipeline/M36_Norms', 'M36_Pipeline_CE/SEED_Train_Norms']
models = ['Model36', 'DINO_DN121', 'model36_large_image', 'EffNetV2S', 'DINO_EffNetV2S', 'EffNetV2S_large_image',
          'RN50', 'DINO_RN50', 'RN50_large_image']
# models = ['Model36', 'DINO_DN121', 'EffNetV2S', 'DINO_EffNetV2S']
# inference_dir = 'inference_test2'
inference_dir = 'inference_test2'
classes = ['3']  # Extend this list with actual class directories
algorithms = ['lac']
alphas = ['alpha_010']
images = [
    'roc_auc_whole_dataset.png',
    'before_removing_any_images_confusion_matrix.png',
    'gt_vs_pred_class_w_avg_pred_set_length_confusion_matrix.png',
    'filtered_by_floor_average_roc_auc_class_1.png',
    'after_removal_by_floor_average_class_1_confusion_matrix.png',
    'Pred Size Vs CoV_boxplot.png',
]

# Function to create PowerPoint slides for each combination
def create_presentation(algorithm, alpha, files, save_path):
    prs = Presentation()
    for file_info in files:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        title = slide.shapes.title
        path_parts = file_info['path'].split('/')
        pipeline1 = path_parts[-7]
        pipeline2 = path_parts[-6]
        model = path_parts[-5]
        num_classes = path_parts[-3]
        avg_pred_set_length = file_info['avg_pred_set_length']
        title.text = f"{pipeline1} - {pipeline2} - {model} - {num_classes} Classes - {algorithm} - {alpha} - Avg Pred Set Length: {avg_pred_set_length}"
        
        # Adjust the font size of the title
        title.text_frame.paragraphs[0].font.size = Pt(14)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Define new image positions for 6 images (3 per row)
        image_positions = [
            (Inches(0.5), Inches(1.5)),  # First image
            (Inches(3.5), Inches(1.5)),  # Second image
            (Inches(6.5), Inches(1.5)),  # Third image
            (Inches(0.5), Inches(4.0)),  # Fourth image
            (Inches(3.5), Inches(4.0)),  # Fifth image
            (Inches(6.5), Inches(4.0))   # Sixth image
        ]
        image_size = (Inches(2.5), Inches(2.5))  # Adjust image size to fit 3x2 layout
        
        for i, (img_path, (left, top)) in enumerate(zip(file_info['images'][:6], image_positions)):
            if os.path.exists(img_path):
                slide.shapes.add_picture(img_path, left, top, width=image_size[0], height=image_size[1])
            else:
                print(f"Image not found: {img_path}")

    prs.save(os.path.join(save_path, f"{algorithm}_{alpha}_top_models_test2_team_presentation_w_large_images.pptx"))

# Function to read the average prediction set length from the CSV file
def read_avg_pred_set_length(csv_file_path):
    with open(csv_file_path, 'r') as csv_file:
        reader = csv.DictReader(csv_file)
        for row in reader:
            return float(row['avg_pred_set_length_whole_df'])
    return None

# Function to gather image paths and CSV data
def gather_files(root_dir, subdirectories, models, inference_dir, classes, algorithms, alphas, images):
    files = {alg: {alpha: [] for alpha in alphas} for alg in algorithms}
    for sub in subdirectories:
        for model in models:
            for num_classes in classes:
                for alg in algorithms:
                    for alpha in alphas:
                        dir_path = os.path.join(root_dir, sub, model, inference_dir, num_classes, alg, alpha)
                        csv_file_path = os.path.join(dir_path, 'analysis_alpha_010')
                        if os.path.exists(dir_path) and os.path.exists(csv_file_path):
                            avg_pred_set_length = read_avg_pred_set_length(csv_file_path)
                            avg_pred_set_length = np.round(avg_pred_set_length, 3)  # Round the value
                            file_info = {
                                'path': dir_path,
                                'avg_pred_set_length': avg_pred_set_length,
                                'images': [os.path.join(dir_path, img) for img in images]
                            }
                            files[alg][alpha].append(file_info)
                        else:
                            print(f"Directory or CSV file not found: {dir_path}")
    return files

if __name__ == "__main__":

    save_path = '/sddata/projects/Conformal_Uncertainty_Quantification/Comparison_SL_SSL_PAVE/analysis/powerpoints'
    # Gather files
    files = gather_files(root_dir, subdirectories, models, inference_dir, classes, algorithms, alphas, images)

    # Create presentations
    for alg in algorithms:
        for alpha in alphas:
            create_presentation(alg, alpha, files[alg][alpha], save_path)
