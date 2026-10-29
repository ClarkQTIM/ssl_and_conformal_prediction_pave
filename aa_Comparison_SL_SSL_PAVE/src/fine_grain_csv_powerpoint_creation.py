import os
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Define the root directory and the combinations of interest
root_dir = '/sddata/projects/Conformal_Uncertainty_Quantification/Comparison_SL_SSL_PAVE/analysis/test2/SSL'
subdirectories = ['M36_Pipeline/M36_Norms']
models = ['Model36', 'DINO_DN121', 'model36_large_image', 'EffNetV2S', 'DINO_EffNetV2S', 'EffNetV2S_large_image',
          'RN50', 'DINO_RN50', 'RN50_large_image']
inference_dir = 'inference_test2'
classes = ['3']  # Extend this list with actual class directories
algorithms = ['lac']
alphas = ['alpha_010']
images = [
    'before_removing_any_images_fine_grain_props.csv',
    'after_removal_by_floor_average_class_1_fine_grain_props.csv',
]

# Function to create PowerPoint slides for each combination
def create_presentation(files, save_path):
    prs = Presentation()
    for file_info in files:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        title = slide.shapes.title
        path_parts = file_info['path'].split(os.sep)
        
        # Extract relevant path information
        pipeline1 = path_parts[-8]
        pipeline2 = path_parts[-7]
        model = path_parts[-6]
        num_classes = path_parts[-4]
        alg = path_parts[-3]
        alpha = path_parts[-2]
        filename = os.path.basename(file_info['path'])
        
        title.text = f"{pipeline1} - {pipeline2} - {model} - {num_classes} Classes - {alg} - {alpha} - {filename}"
        
        # Adjust the font size of the title
        title.text_frame.paragraphs[0].font.size = Pt(14)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        left = Inches(1)
        top = Inches(1.5)
        width = Inches(8.5)
        height = Inches(5)

        if file_info['path'].endswith('.csv') and os.path.exists(file_info['path']):
            # Add table for CSV data
            with open(file_info['path'], 'r') as csv_file:
                reader = csv.reader(csv_file)
                table_data = list(reader)

            # Define table position and size
            rows, cols = len(table_data), len(table_data[0])
            table = slide.shapes.add_table(rows, cols, left, top, width, height).table

            # Set column widths
            col_width = int(width / cols)
            for col in range(cols):
                table.columns[col].width = col_width

            # Fill the table with CSV data
            for row in range(rows):
                for col in range(cols):
                    cell = table.cell(row, col)
                    cell.text = table_data[row][col]
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(10)  # Set font size

    prs.save(save_path)
    print(f"Presentation saved to: {save_path}")

# Function to gather file paths
def gather_files(root_dir, subdirectories, models, inference_dir, classes, algorithms, alphas, images):
    files = []
    for sub in subdirectories:
        for model in models:
            for num_classes in classes:
                for alg in algorithms:
                    for alpha in alphas:
                        dir_path = os.path.join(root_dir, sub, model, inference_dir, num_classes, alg, alpha)
                        for img in images:
                            img_path = os.path.join(dir_path, img)
                            if os.path.exists(img_path):
                                file_info = {
                                    'path': img_path
                                }
                                files.append(file_info)
                            else:
                                print(f"File not found: {img_path}")
    return files

if __name__ == "__main__":
    save_path = '/sddata/projects/Conformal_Uncertainty_Quantification/Comparison_SL_SSL_PAVE/analysis/powerpoints/ssl_fine_grain_props_presentation.pptx'
    # Gather files
    files = gather_files(root_dir, subdirectories, models, inference_dir, classes, algorithms, alphas, images)

    # Create presentation
    create_presentation(files, save_path)
